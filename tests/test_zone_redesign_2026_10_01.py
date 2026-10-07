"""2026-10-01 Zone 탭 전면 재설계 — 라운드 D(append+삭제)/E(레시피+수동편집)/
F(결과 분석 테이블) 회귀 테스트.

스펙: docs/specs/zone-tab-redesign-2026-10-01.md
"""
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor, QImage, QPixmap
from PyQt6.QtWidgets import QApplication

from app.widgets.inference_image_list import InferenceImageList
from app.widgets.zone_canvas import ZoneCanvas
from app.widgets.zone_recipe_dialog import ZoneRecipeDialog
from app.widgets.zone_batch_result_dialog import ZoneBatchResultDialog
from app.core import zone_recipe_store as recipe_store
from app.core.zone_metrics import scale_circles, max_blob_pixels_by_zone, ZoneBlobStat
from app.tabs.zone_analysis_tab import ZoneAnalysisTab

_APP = QApplication.instance() or QApplication(sys.argv)


def _stat(zone_name: str, blob_id: int, pixel_count: int) -> ZoneBlobStat:
    return ZoneBlobStat(
        zone_name=zone_name, blob_id=blob_id, pixel_count=pixel_count,
        ai_score=0.9, centroid_x=1.0, centroid_y=1.0,
        bbox_x=0, bbox_y=0, bbox_w=1, bbox_h=1,
    )


def _save_png(path: Path) -> None:
    assert QImage(4, 4, QImage.Format.Format_RGB32).save(str(path))


# ── 라운드 D: append + 개별 삭제 ────────────────────────────────────────────

def test_load_files_append_merges_and_dedupes(tmp_path: Path) -> None:
    p1, p2, p3 = tmp_path / "a.png", tmp_path / "b.png", tmp_path / "c.png"
    for p in (p1, p2, p3):
        _save_png(p)

    lst = InferenceImageList()
    lst.load_files([p1, p2])
    assert set(lst.paths()) == {p1, p2}

    lst.load_files([p2, p3], append=True)   # p2 중복 — set 연산으로 제거
    assert set(lst.paths()) == {p1, p2, p3}

    lst.load_files([p1], append=False)   # 기본값 False는 기존처럼 전체 교체
    assert set(lst.paths()) == {p1}


def test_load_folder_append_root_mismatch_falls_back_to_flat_tree(tmp_path: Path) -> None:
    root_a = tmp_path / "root_a"
    root_b = tmp_path / "root_b"
    root_a.mkdir()
    root_b.mkdir()
    pa, pb = root_a / "a.png", root_b / "b.png"
    _save_png(pa)
    _save_png(pb)

    lst = InferenceImageList()
    lst.load_folder(root_a)
    assert lst._root == root_a

    lst.load_folder(root_b, append=True)   # 루트 불일치 -> 평탄 그룹 트리로 폴백
    assert lst._root is None
    assert set(lst.paths()) == {pa, pb}


def test_remove_selected_keeps_files_emits_signal_and_clears_status(tmp_path: Path) -> None:
    p1, p2 = tmp_path / "a.png", tmp_path / "b.png"
    _save_png(p1)
    _save_png(p2)

    lst = InferenceImageList()
    lst.load_files([p1, p2])
    lst.set_item_status(p1, "done", badge="완료")

    removed = []
    lst.images_removed.connect(lambda paths: removed.extend(paths))
    lst._tree.setCurrentItem(lst._path_to_item[p1])
    lst._remove_selected()

    assert removed == [p1]
    assert set(lst.paths()) == {p2}
    assert p1 not in lst._status
    assert p1.exists(), "목록에서 제거해도 원본 파일은 보존돼야 한다"


def test_zone_tab_images_removed_clears_results_and_current_canvas(tmp_path: Path) -> None:
    p1, p2 = tmp_path / "a.png", tmp_path / "b.png"
    _save_png(p1)
    _save_png(p2)
    tab = ZoneAnalysisTab()
    try:
        tab._img_list.load_files([p1, p2])
        tab._results[p1] = object()
        tab._results[p2] = object()
        tab._image_path = p1

        tab._on_images_removed([p1])

        assert p1 not in tab._results
        assert p2 in tab._results
        assert tab._image_path is None
    finally:
        tab.close()


# ── 라운드 E: 수동 원 편집(방향키/휠/정렬) ───────────────────────────────────

def test_align_centers_averages_without_changing_radius() -> None:
    canvas = ZoneCanvas()
    canvas.set_image_size(100, 100)
    canvas.set_circles([(0.0, 0.0, 5.0), (10.0, 20.0, 9.0)])

    canvas.align_centers()

    circles = canvas.get_circles()
    assert len(circles) == 2
    for cx, cy, r, _name in circles:
        assert abs(cx - 5.0) < 1e-6 and abs(cy - 10.0) < 1e-6
    assert sorted(r for _, _, r, _name in circles) == [5.0, 9.0]


def test_align_centers_noop_with_fewer_than_two_circles() -> None:
    canvas = ZoneCanvas()
    canvas.set_image_size(100, 100)
    canvas.set_circles([(1.0, 2.0, 3.0)])
    before = canvas.get_circles()

    canvas.align_centers()

    assert canvas.get_circles() == before


def test_alt_wheel_resizes_selected_circle_plain_wheel_always_zooms() -> None:
    """2026-10-03#6: Alt+휠=지름 조절, 일반 휠(원 선택 여부 무관)=항상 화면 줌."""
    canvas = ZoneCanvas()
    canvas.resize(200, 200)
    canvas.set_image_size(100, 100)
    canvas.set_pixmap(QPixmap(100, 100))
    canvas.set_circles([(50.0, 50.0, 10.0)])
    circle_id = canvas.circles_with_ids()[0][0]
    canvas.select_circle(circle_id)

    from PyQt6.QtGui import QWheelEvent
    from PyQt6.QtCore import QPointF, QPoint

    def make_event(modifiers: Qt.KeyboardModifier) -> QWheelEvent:
        return QWheelEvent(
            QPointF(50, 50), QPointF(50, 50), QPoint(0, 0), QPoint(0, 120),
            Qt.MouseButton.NoButton, modifiers,
            Qt.ScrollPhase.NoScrollPhase, False,
        )

    zoom_before = canvas._zoom
    undo_depth_before = len(canvas._undo_stack)
    canvas.wheelEvent(make_event(Qt.KeyboardModifier.NoModifier))
    _, _, r, _name = canvas.get_circles()[0]
    assert r == 10.0   # 일반 휠은 지름을 건드리지 않는다
    assert canvas._zoom != zoom_before   # 대신 화면 줌이 동작한다
    assert len(canvas._undo_stack) == undo_depth_before   # 줌은 undo 스택에 쌓이지 않는다

    canvas.wheelEvent(make_event(Qt.KeyboardModifier.AltModifier))
    _, _, r, _name = canvas.get_circles()[0]
    assert r > 10.0   # Alt+휠 위로 스크롤 -> 반지름(지름) 증가
    assert canvas.can_undo()


def test_arrow_key_moves_selected_circle_with_gesture_debounced_undo() -> None:
    canvas = ZoneCanvas()
    canvas.resize(200, 200)
    canvas.set_image_size(100, 100)
    canvas.set_pixmap(QPixmap(100, 100))
    canvas.set_circles([(50.0, 50.0, 10.0)])
    circle_id = canvas.circles_with_ids()[0][0]
    canvas.select_circle(circle_id)
    undo_depth_before = len(canvas._undo_stack)

    from PyQt6.QtGui import QKeyEvent
    from PyQt6.QtCore import QEvent

    for _ in range(3):   # 연타 -> 디바운스 때문에 undo 스택엔 1개만 쌓여야 함
        ev = QKeyEvent(QEvent.Type.KeyPress, Qt.Key.Key_Right, Qt.KeyboardModifier.NoModifier)
        canvas.keyPressEvent(ev)

    cx, cy, _, _name = canvas.get_circles()[0]
    assert cx == 53.0   # 1px * 3회
    assert len(canvas._undo_stack) == undo_depth_before + 1   # 제스처 디바운스 -> 1개만

    canvas.undo()
    cx2, _, _, _name = canvas.get_circles()[0]
    assert cx2 == 50.0   # 디바운스된 제스처 전체가 undo 1회로 복원


# ── 라운드 E: 레시피 저장/불러오기 + 다이얼로그 게이트 ───────────────────────

def test_recipe_store_round_trip_and_scale(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    circles = [(10.0, 10.0, 5.0), (10.0, 10.0, 8.0)]
    path = recipe_store.save_recipe("테스트 레시피", circles, (100, 100))

    recipes = recipe_store.list_recipes()
    assert recipes == [path]

    loaded = recipe_store.load_recipe(path)
    assert loaded["name"] == "테스트 레시피"
    assert loaded["circles"] == circles

    scaled = scale_circles(loaded["circles"], loaded["ref_size"], (200, 200))
    assert scaled == [(20.0, 20.0, 10.0), (20.0, 20.0, 16.0)]


def test_recipe_dialog_apply_button_gated_on_circle_count(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)   # 레시피 디렉터리를 격리(빈 상태로 시작)
    pixmap = QPixmap(50, 50)
    pixmap.fill(QColor("white"))
    dialog = ZoneRecipeDialog(pixmap, (50, 50))
    try:
        assert not dialog._btn_apply.isEnabled()   # 레시피 없음 -> 빈 캔버스로 시작

        dialog._canvas.set_circles([(25.0, 25.0, 10.0)])
        assert dialog._btn_apply.isEnabled()

        dialog._canvas.clear_circles()
        assert not dialog._btn_apply.isEnabled()

        dialog._canvas.set_circles([(25.0, 25.0, 10.0)])
        dialog._on_apply()
        assert dialog.result_circles() == [(25.0, 25.0, 10.0, None)]
        assert dialog.result_ref_size() == (50, 50)
    finally:
        dialog.close()


def test_recipe_dialog_autoloads_most_recent_recipe(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    recipe_store.save_recipe("최근 레시피", [(1.0, 1.0, 1.0)], (10, 10))
    pixmap = QPixmap(20, 20)
    dialog = ZoneRecipeDialog(pixmap, (20, 20))
    try:
        assert dialog._canvas.get_circles() == [(2.0, 2.0, 2.0, None)]   # (10,10)->(20,20) 2배 스케일
        assert dialog._btn_apply.isEnabled()
    finally:
        dialog.close()


# ── 라운드 E: 활성화 조건 분리(7-2) — 추론 전에도 자동검출 가능 ──────────────

def test_detect_button_enabled_right_after_image_load_before_inference(tmp_path: Path) -> None:
    img_path = tmp_path / "img.png"
    assert QImage(20, 20, QImage.Format.Format_RGB32).save(str(img_path))
    tab = ZoneAnalysisTab()
    try:
        assert not tab._btn_detect.isEnabled()   # 이미지 로드 전

        tab._on_list_image_selected(img_path)

        assert tab._btn_detect.isEnabled()   # 추론 없이도 바로 활성화(7-2)
        assert tab._act_circle.isEnabled()
        for action in (tab._act_brush_draw, tab._act_brush_erase, tab._act_blob_delete):
            assert not action.isEnabled()   # 블랍 기반 편집은 여전히 추론 후에만
    finally:
        tab.close()


def test_recipe_button_visibility_follows_batch_mode() -> None:
    tab = ZoneAnalysisTab()
    try:
        assert tab._mode_combo.currentData() == "apply_all"
        assert not tab._btn_recipe.isHidden()

        per_image_idx = [tab._mode_combo.itemData(i) for i in range(tab._mode_combo.count())].index("per_image")
        tab._mode_combo.setCurrentIndex(per_image_idx)
        assert tab._btn_recipe.isHidden()
    finally:
        tab.close()


# ── 라운드 F: 결과 분석 테이블(최대 blob 픽셀수 + 그룹화 + 필터 + 클립보드) ──

def test_max_blob_pixels_by_zone_picks_max_per_image_and_zone() -> None:
    blob_rows = [
        ("a.png", _stat("중심부", 1, 100)),
        ("a.png", _stat("중심부", 2, 300)),
        ("a.png", _stat("바깥쪽", 1, 50)),
        ("b.png", _stat("중심부", 1, 10)),
    ]
    result = max_blob_pixels_by_zone(blob_rows)
    assert result == {
        ("a.png", "중심부"): 300, ("a.png", "바깥쪽"): 50, ("b.png", "중심부"): 10,
    }


def test_export_excel_zones_sheet_has_max_blob_column(tmp_path: Path) -> None:
    from app.core.zone_metrics import export_zone_percentages_to_excel
    from openpyxl import load_workbook

    rows = [("a.png", "중심부", 80.0), ("a.png", "바깥쪽", 5.0)]
    blob_rows = [("a.png", _stat("중심부", 1, 123))]
    out = tmp_path / "out.xlsx"

    export_zone_percentages_to_excel(rows, out, blob_rows)

    wb = load_workbook(out)
    ws = wb["zones"]
    assert [c.value for c in ws[1]] == ["이미지파일명", "존이름", "타겟비율(%)", "최대 blob 픽셀수"]
    assert ws[2][3].value == 123   # 중심부 blob 있음
    assert ws[3][3].value == 0     # 바깥쪽은 blob 없음 -> 0


def test_long_tab_groups_same_image_rows_via_span_and_sorts_by_image() -> None:
    rows = [("b.png", "중심부", 10.0), ("a.png", "중심부", 20.0), ("a.png", "바깥쪽", 5.0)]
    blob_rows: list = []
    dialog = ZoneBatchResultDialog(rows, blob_rows)
    try:
        # 이미지명 기준 정렬 -> a.png(2행) 먼저, b.png(1행) 나중
        assert [r[0] for r in dialog._long_rows_sorted] == ["a.png", "a.png", "b.png"]
        assert dialog._long_table.rowSpan(0, 0) == 2   # a.png 2행 그룹화
        assert dialog._long_table.rowSpan(2, 0) == 1   # b.png 1행은 병합 없음
    finally:
        dialog.close()


def test_filter_bar_search_and_zone_toggle_affect_only_visible_rows() -> None:
    rows = [("a.png", "중심부", 10.0), ("b.png", "중심부", 20.0), ("b.png", "바깥쪽", 5.0)]
    dialog = ZoneBatchResultDialog(rows, [])
    try:
        # 초기 상태 — 전부 표시
        assert dialog._lbl_filter_count.text() == "3 / 3"

        dialog._search_edit.setText("b.png")
        assert dialog._lbl_filter_count.text() == "2 / 3"
        for r, (img, _zone, _pct) in enumerate(dialog._long_rows_sorted):
            assert dialog._long_table.isRowHidden(r) == (img != "b.png")

        dialog._on_reset_filter()
        assert dialog._lbl_filter_count.text() == "3 / 3"

        dialog._zone_buttons["바깥쪽"].setChecked(False)
        assert dialog._lbl_filter_count.text() == "2 / 3"

        # 내보내기 데이터는 필터와 무관하게 항상 전체
        assert len(dialog._rows) == 3
    finally:
        dialog.close()


def test_clipboard_copy_includes_max_blob_column(monkeypatch) -> None:
    rows = [("a.png", "중심부", 80.0)]
    blob_rows = [("a.png", _stat("중심부", 1, 999))]
    dialog = ZoneBatchResultDialog(rows, blob_rows)
    try:
        dialog._on_copy_clipboard()
        text = QApplication.clipboard().text()
        assert "이미지\t존\t타겟 비율(%)\t최대 blob 픽셀수" in text
        assert "a.png\t중심부\t80.00\t999" in text
    finally:
        dialog.close()


def test_export_single_reuses_batch_result_dialog(tmp_path: Path, monkeypatch) -> None:
    img_path = tmp_path / "img.png"
    assert QImage(20, 20, QImage.Format.Format_RGB32).save(str(img_path))
    tab = ZoneAnalysisTab()
    opened = []

    class _Dlg:
        def __init__(self, rows, blob_rows, parent=None):
            opened.append((rows, blob_rows))

        def exec(self):
            return 0

    import app.tabs.zone_analysis_tab as module
    monkeypatch.setattr(module, "ZoneBatchResultDialog", _Dlg)
    # BUG-040 B: _compute_zone_percentages/_compute_zone_blob_rows가 _current_zones()로
    # 공유 계산한 결과를 받는 선택적 zones 인자를 받도록 바뀌어 호출부가 위치 인자를 넘긴다.
    monkeypatch.setattr(tab, "_compute_zone_percentages", lambda zones=None: [("중심부", 50.0)])
    monkeypatch.setattr(tab, "_compute_zone_blob_rows", lambda zones=None: [])
    tab._image_path = img_path

    try:
        tab._on_export_single()
        assert len(opened) == 1
        rows, blob_rows = opened[0]
        assert rows == [("img.png", "중심부", 50.0)]
        assert blob_rows == []
    finally:
        tab.close()


if __name__ == "__main__":
    import pytest as _pytest
    raise SystemExit(_pytest.main([__file__, "-q"]))
