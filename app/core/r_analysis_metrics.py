"""R 분석(R Analysis) 탭 — 임계값 이진화 + contour 기반 면적 검출, 셀 단위 ratio/판정.

스펙: docs/specs/r-analysis-tab-2026-10-26.md 요구사항 1/6/8. Qt 의존성 없음(core 규칙).
"셀(cell)"은 이미지 여러 장을 묶는 단위 — R_cell = sum(area) / denominator를 셀당
1번만 계산해 스펙 상/하한과 비교한다(이미지별 개별 R/판정은 없음, 요구사항 6).

`zone_metrics.disk_mask`/`zone_metrics.apply_manual_strokes`를 재사용한다 — 마스크
보정(브러시 그리기/지우기) 수학은 존 분석 탭과 동일(요구사항 7, 원↔임의 블롭 모양
차이는 `r_roi_canvas.py`가 복사 포팅으로 흡수).
"""
from __future__ import annotations

from pathlib import Path

import cv2
import numpy as np

from app.core.zone_metrics import disk_mask, apply_manual_strokes   # noqa: F401 (재사용 — 공개 재노출)

# ponytail: 1px 수준의 이진화 노이즈를 거르는 최소 면적 — 값 자체를 조정 가능한
# 설정으로 노출하지 않는다(요청받지 않은 옵션, YAGNI). 실측상 노이즈가 더 커지면
# 그때 매개변수화.
_MIN_CONTOUR_AREA = 4.0


def detect_contour_mask(
    img: np.ndarray, roi: tuple[int, int, int, int], threshold: int,
) -> np.ndarray:
    """ROI로 crop한 뒤 이진화 -> `cv2.findContours`(RETR_EXTERNAL)로 찾은 모든
    유효 contour(노이즈 최소면적 이상)를 합쳐 원본 이미지 크기의 bool 마스크로
    반환한다. 단일 최대 contour만 쓰지 않는다 — 검사 대상이 여러 블롭으로
    나뉠 수 있기 때문(요구사항 1).

    # ponytail: 교체 지점 — 추후 AI 모델 기반으로 바꿀 때 이 함수 시그니처
    # (img, roi, threshold)->mask 유지
    """
    h, w = img.shape[:2]
    mask = np.zeros((h, w), dtype=bool)
    x, y, rw, rh = roi
    x0, y0 = max(0, int(x)), max(0, int(y))
    x1, y1 = min(w, x0 + int(rw)), min(h, y0 + int(rh))
    if x1 <= x0 or y1 <= y0:
        return mask

    crop = img[y0:y1, x0:x1]
    gray = cv2.cvtColor(crop, cv2.COLOR_BGR2GRAY) if crop.ndim == 3 else crop
    _, binary = cv2.threshold(gray, threshold, 255, cv2.THRESH_BINARY)
    contours, _ = cv2.findContours(binary, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

    local = np.zeros(binary.shape, dtype=np.uint8)
    for c in contours:
        if cv2.contourArea(c) >= _MIN_CONTOUR_AREA:
            cv2.drawContours(local, [c], -1, 1, thickness=cv2.FILLED)
    mask[y0:y1, x0:x1] = local.astype(bool)
    return mask


def contour_area(mask: np.ndarray) -> int:
    return int(mask.sum())


def compute_cell_ratio(sum_area: float, denominator: float) -> float | None:
    if denominator <= 0:
        return None
    return sum_area / denominator


def pass_fail(ratio: float | None, spec_lower: float, spec_upper: float) -> bool | None:
    if ratio is None:
        return None
    return spec_lower <= ratio <= spec_upper


def cell_rollup(
    per_image_areas: list[int], denominator: float, spec_lower: float, spec_upper: float,
) -> dict:
    """셀 내 이미지별 면적 목록 -> 셀 단위 합계/ratio/판정 1건(요구사항 6)."""
    sum_area = sum(per_image_areas)
    ratio = compute_cell_ratio(sum_area, denominator)
    return {
        "sum_area": int(sum_area),
        "image_count": len(per_image_areas),
        "ratio": ratio,
        "pass_fail": pass_fail(ratio, spec_lower, spec_upper),
    }


def export_r_results_to_excel(cells: dict, out_path: Path) -> None:
    """cells: {cell_id: {"sum_area": int, "image_count": int, "ratio": float|None,
    "pass_fail": bool|None, "images": [(image_name, area), ...]}}

    시트 "cells"(셀별 요약 1행) + 시트 "images"(이미지별 면적, 참고용 — 요구사항 8).
    `zone_metrics.export_zone_percentages_to_excel`과 동일한 openpyxl 패턴(헤더 볼드).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "cells"
    ws.append(["cell_id", "image_count", "sum_area", "ratio", "pass_fail"])
    for header_cell in ws[1]:
        header_cell.font = Font(bold=True)
    for cell_id, data in cells.items():
        ratio = data.get("ratio")
        pf = data.get("pass_fail")
        ws.append([
            cell_id,
            data.get("image_count", 0),
            data.get("sum_area", 0),
            round(ratio, 4) if ratio is not None else "N/A",
            "PASS" if pf else ("FAIL" if pf is False else "N/A"),
        ])

    ws2 = wb.create_sheet("images")
    ws2.append(["cell_id", "image_name", "area"])
    for header_cell in ws2[1]:
        header_cell.font = Font(bold=True)
    for cell_id, data in cells.items():
        for image_name, area in data.get("images", []):
            ws2.append([cell_id, image_name, area])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))


if __name__ == "__main__":
    # ── self-check: 합성 이미지에 알려진 면적(15행 x 20열 = 300px)의 사각형 ──────
    img = np.zeros((50, 50), dtype=np.uint8)
    img[10:25, 10:30] = 255
    mask = detect_contour_mask(img, (0, 0, 50, 50), 127)
    area = contour_area(mask)
    assert area == 300, f"300 예상, 실측 {area}"

    # ROI 밖은 무시
    mask_roi = detect_contour_mask(img, (0, 0, 5, 5), 127)
    assert contour_area(mask_roi) == 0, "ROI 밖 사각형은 검출되지 않아야 함"

    # 1px 노이즈는 최소면적 미달로 제외
    img_noisy = img.copy()
    img_noisy[0, 0] = 255
    mask_noisy = detect_contour_mask(img_noisy, (0, 0, 50, 50), 127)
    assert contour_area(mask_noisy) == 300, "1px 노이즈는 결과 면적에 포함되면 안 됨"

    # 떨어진 블롭 2개 -> 합쳐서 검출(단일 최대 contour만 쓰지 않음, 요구사항 1)
    img2 = np.zeros((50, 50), dtype=np.uint8)
    img2[5:10, 5:10] = 255     # 25px
    img2[30:40, 30:35] = 255   # 50px
    mask2 = detect_contour_mask(img2, (0, 0, 50, 50), 127)
    assert contour_area(mask2) == 75, f"두 블롭 합계 75 예상, 실측 {contour_area(mask2)}"

    # cell_rollup / pass_fail 산술
    rollup = cell_rollup([100, 150, 50], denominator=300.0, spec_lower=0.9, spec_upper=1.1)
    assert rollup == {"sum_area": 300, "image_count": 3, "ratio": 1.0, "pass_fail": True}, rollup

    rollup_fail = cell_rollup([10], denominator=300.0, spec_lower=0.9, spec_upper=1.1)
    assert rollup_fail["pass_fail"] is False

    rollup_none = cell_rollup([10], denominator=0.0, spec_lower=0.9, spec_upper=1.1)
    assert rollup_none["ratio"] is None and rollup_none["pass_fail"] is None

    assert compute_cell_ratio(10, 0) is None
    assert compute_cell_ratio(10, 0.0) is None
    assert compute_cell_ratio(10, -5) is None
    assert pass_fail(None, 0.9, 1.1) is None
    assert pass_fail(1.0, 0.9, 1.1) is True
    assert pass_fail(0.5, 0.9, 1.1) is False

    # zone_metrics 재사용 확인(재구현 아님)
    assert disk_mask(2, 2, 1, (5, 5)).sum() == 5
    base = np.zeros((9, 9), dtype=bool)
    edited = apply_manual_strokes(base, [(True, [(4.0, 4.0, 2.0)])])
    assert edited[4, 4]

    # export_r_results_to_excel — 왕복(파일 생성 + 시트/행 수 확인)
    import tempfile
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "r_results.xlsx"
        export_r_results_to_excel(
            {"CELL-A": {**rollup, "images": [("a.jpg", 100), ("b.jpg", 150), ("c.jpg", 50)]}},
            out,
        )
        wb = load_workbook(out)
        assert wb.sheetnames == ["cells", "images"]
        assert wb["cells"].max_row == 2   # 헤더 + 1행
        assert wb["images"].max_row == 4  # 헤더 + 3행

    print("r_analysis_metrics self-check OK")
